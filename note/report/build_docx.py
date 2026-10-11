# -*- coding: utf-8 -*-
"""Assemble the report markdown chapters into one .docx (python-docx).

Covers the markdown subset used by the report: headings, paragraphs with **bold** and
`inline code`, bullet/numbered lists, pipe tables, fenced code blocks, images and
> blockquotes. Run from the repo root:
    python note/report/build_docx.py
"""
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

CHAPTERS = [
    'note/report/01-软件介绍.md',
    'note/report/02-架构设计.md',
    'note/report/03-关键模块详细设计.md',
    'note/report/04-软件维护.md',
]
COVER = 'note/report/开源软件阅读维护报告.md'
OUT = 'note/report/开源软件阅读维护报告.docx'

CN = '微软雅黑'
MONO = 'Consolas'


def set_cn(style, font=CN, size=None):
    style.font.name = 'Calibri'
    style.element.rPr.rFonts.set(qn('w:eastAsia'), font)
    if size:
        style.font.size = Pt(size)


def add_runs(par, text):
    """Add runs to a paragraph, honouring **bold** and `code` spans."""
    for part in re.split(r'(\*\*.+?\*\*|`[^`]+`)', text):
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            r = par.add_run(part[2:-2])
            r.bold = True
        elif part.startswith('`') and part.endswith('`'):
            r = par.add_run(part[1:-1])
            r.font.name = MONO
            r._element.rPr.rFonts.set(qn('w:eastAsia'), MONO)
            r.font.color.rgb = RGBColor(0x8B, 0x2F, 0x00)
        else:
            par.add_run(part)


def add_table(doc, rows):
    cells = [[c.strip() for c in r.strip().strip('|').split('|')] for r in rows]
    cells = [r for r in cells if not all(re.fullmatch(r':?-{2,}:?', c or '---') for c in r)]
    if not cells:
        return
    t = doc.add_table(rows=len(cells), cols=len(cells[0]))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(cells):
        for j, val in enumerate(row):
            if j < len(t.rows[i].cells):
                p = t.rows[i].cells[j].paragraphs[0]
                add_runs(p, val)
                for r in p.runs:
                    r.font.size = Pt(9)
                    if i == 0:
                        r.bold = True


def add_code(doc, lines):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.6)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    shade = p._element.get_or_add_pPr()
    for i, line in enumerate(lines):
        r = p.add_run(line)
        r.font.name = MONO
        r._element.rPr.rFonts.set(qn('w:eastAsia'), MONO)
        r.font.size = Pt(8.5)
        if i < len(lines) - 1:
            r.add_break()
    shd = shade.makeelement(qn('w:shd'), {qn('w:val'): 'clear', qn('w:fill'): 'F2F2F2'})
    shade.append(shd)


def render(doc, path, top_heading=1):
    import os
    base = os.path.dirname(path)
    in_code, code_buf = False, []
    table_buf = []
    for raw in open(path, encoding='utf-8').read().splitlines():
        line = raw.rstrip()
        if line.startswith('```'):
            if in_code:
                add_code(doc, code_buf)
                code_buf = []
            in_code = not in_code
            continue
        if in_code:
            code_buf.append(raw)
            continue
        if line.startswith('|'):
            table_buf.append(line)
            continue
        if table_buf:
            add_table(doc, table_buf)
            table_buf = []
        if not line.strip():
            continue
        m = re.match(r'^(#{1,4})\s+(.*)', line)
        if m:
            level = len(m.group(1)) + top_heading - 1
            level = min(level, 4)
            h = doc.add_heading('', level=level)
            add_runs(h, m.group(2))
            continue
        if line.startswith('!['):
            img = re.match(r'^!\[.*\]\((.+)\)', line)
            if img:
                doc.add_picture(os.path.join(base, img.group(1)), width=Cm(15.5))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            continue
        if line.startswith('> '):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            add_runs(p, line[2:])
            for r in p.runs:
                r.italic = True
                r.font.color.rgb = RGBColor(0x60, 0x60, 0x60)
            continue
        if re.match(r'^[-*]\s+', line):
            p = doc.add_paragraph(style='List Bullet')
            add_runs(p, re.sub(r'^[-*]\s+', '', line))
            continue
        if re.match(r'^\d+\.\s+', line):
            p = doc.add_paragraph(style='List Number')
            add_runs(p, re.sub(r'^\d+\.\s+', '', line))
            continue
        p = doc.add_paragraph()
        add_runs(p, line)
    if table_buf:
        add_table(doc, table_buf)


doc = Document()
set_cn(doc.styles['Normal'], size=10.5)
doc.styles['Normal'].paragraph_format.space_after = Pt(6)
for name, size in [('Heading 1', 18), ('Heading 2', 14), ('Heading 3', 12), ('Heading 4', 11)]:
    set_cn(doc.styles[name], size=size)
    doc.styles[name].font.color.rgb = RGBColor(0x1F, 0x3F, 0x66)

# Cover from the master file, then the four chapters (their H1 becomes a chapter heading).
render(doc, COVER, top_heading=0)
for ch in CHAPTERS:
    doc.add_page_break()
    render(doc, ch)

doc.save(OUT)
print('wrote', OUT)
